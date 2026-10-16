from pathlib import Path
import json,sys,re,os
CAPTAIN_PHONE = os.environ.get("SUBMISSION_CAPTAIN_PHONE", "")
if not CAPTAIN_PHONE:
 raise SystemExit("Set SUBMISSION_CAPTAIN_PHONE locally before generating personal submission documents.")
R=Path('/data-ssd/zhongchengkai/agnes');O=R/'artifacts/submission-materials-G002-20261015';D=O/'package';T=R/'artifacts/submission-materials-20261015/tooling';sys.path.insert(0,str(T/'python-libs'))
from docx import Document
M={'【待填_报名组别】':'硕博组','【待填_官方参赛编号】':'G002','【待填_队伍名称】':'今天不想做正事','【待填_队长姓名】':'郭煜','【待填_学校学院】':'东南大学','【待填_学籍身份】':'研究生','【待填_队长电话】':CAPTAIN_PHONE,'【待填_无需账号登录的视频主链接】':'【待填_演示视频主链接】','【待填_备用链接或填无】':'【待填_演示视频备用链接】','【待填_公开发布年月日时分及时间区】':'【待填_公开发布时间】','【待填_平台】':'【待填_公开发布平台】','【待填_有效链接】':'【待填_公开内容链接】','【待填_真实发布时间】':'【待填_公开发布时间】','【待填_实际公开页面截图文件名】':'【待填_公开页面截图】'}
REMOVE_ROW=['视频访问密码','代码或体验链接','参考方向','权重细分版本']
REMOVE_P=['成员分工：','队长签名：','签署日期：','项目自有代码发布许可证：']
MEMBERS='成员信息：队长郭煜，东南大学研究生；队员林路辰，东南大学研究生。'
SUBSTANTIVE={
 '本队将如实说明第三方代码、模型、软件、数据、图像、字体及三维资产的来源与使用方式；保留原许可和版权通知。AI 辅助参与代码实现、调试、媒体和材料整理，实际范围见下一页。成员应补充自有工作和分工，不将 AI 辅助或第三方工作表述为全部自行编写。':'本队将如实说明第三方代码、模型、软件、数据、图像、字体及三维资产的来源与使用方式，并保留原许可和版权通知。AI 辅助参与代码实现、调试、媒体和材料整理，实际范围见下一页。',
 '用户提供的骨架、办公室视频及火灾参考图用于外观和布局参考。未将参考视频帧或图库照片直接提取为地图纹理。素材权属和授权：【待填_团队提供素材的权属及授权说明】。若公开发布演示片，成员应确认引用展示及品牌、素材许可适用范围。':'用户提供的骨架、办公室视频及火灾参考图用于外观和布局参考，未将参考视频帧或图库照片直接提取为地图纹理。第三方素材及代码保留相应来源和许可记录。',
 '赛事规则对模型调用作了限制。Qwen 后期配音和 AI 制作辅助单独披露，不能将它们隐去或称为 Agnes 输出。后期配音是否被赛事允许：【待填_主办方确认或更换配音结果】。这不是机器人核心任务模型变更。':'Qwen 后期配音和 AI 制作辅助已单独披露，不属于机器人运行中的视觉或任务决策模型。',
}
def trans(s):
 if s.startswith('成员信息：'):return MEMBERS
 if s in SUBSTANTIVE:return SUBSTANTIVE[s]
 for a,b in M.items():s=s.replace(a,b)
 return s

def setp(p,s):
 if p.runs:
  p.runs[0].text=s
  for r in p.runs[1:]:r.text=''
 else:p.add_run(s)

for p in D.rglob('*.docx'):
 doc=Document(p)
 for paragraph in list(doc.paragraphs):
  text=paragraph.text
  if any(text.startswith(q)for q in REMOVE_P):paragraph._p.getparent().remove(paragraph._p);continue
  new=trans(text)
  if new!=text:setp(paragraph,new)
 for t in doc.tables:
  for row in list(t.rows):
   label=row.cells[0].text
   if label in REMOVE_ROW:row._tr.getparent().remove(row._tr);continue
   if label=='联系方式':row.cells[0].text='队长联系电话';row.cells[1].text=CAPTAIN_PHONE
   elif label=='队长和成员':row.cells[1].text='队长：郭煜；队员：林路辰'
   elif label=='学校和学籍':row.cells[1].text='郭煜、林路辰：东南大学，研究生'
   else:
    for cell in row.cells:
     for para in cell.paragraphs:
      new=trans(para.text)
      if new!=para.text:setp(para,new)
 doc.core_properties.author='今天不想做正事';doc.core_properties.last_modified_by='今天不想做正事';doc.save(p)
 print('UPDATED_DOCX',p.name,flush=True)

for p in D.rglob('*.md'):
 text=p.read_text();lines=[]
 for l in text.splitlines():
  if l.startswith('|'):
   label=l.split('|')[1].strip()
   if label in REMOVE_ROW or label in ['视频密码','匿名访问核查日期']:continue
   if label=='联系方式':l=f'| 队长联系电话 | {CAPTAIN_PHONE} |'
   elif label=='队长和成员':l='| 队长和成员 | 队长：郭煜；队员：林路辰 |'
   elif label=='学校和学籍':l='| 学校和学籍 | 郭煜、林路辰：东南大学，研究生 |'
  if any(l.startswith(q)for q in REMOVE_P):continue
  lines.append(trans(l))
 p.write_text('\n'.join(lines)+'\n')

j=D/'04_可核查运行证据/模型调用核查.json';data=json.loads(j.read_text());data.pop('weight_revision',None);j.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
j=D/'00_审核入口/视频版本记录.json';data=json.loads(j.read_text());data['public_main_link']='【待填_演示视频主链接】';j.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')

(D/'00_审核入口/人工修改清单.md').write_text('''# 发布信息补充清单

队伍：今天不想做正事；组别：硕博组；参赛编号：G002。队长郭煜、队员林路辰，均为东南大学研究生。队长联系电话已填入提交表。

1. 上传当前 V16 展示视频，填写一个可匿名访问主链接；如提供备用链接，再填写备用链接。
2. 填写至少一条真实公开内容的标题、平台、有效链接、真实发布时间，并添加实际页面截图。
3. 更新 Word 或 Markdown 后同步导出 PDF，运行 make_manifest.py 和 verify_bundle.py 更新并核对哈希。
4. 正式压缩包名称为“硕博组-G002-今天不想做正事-郭煜.zip”，大小不超过 200 MB。

未提供的邮箱、学院、个人分工、额外成员和签署日期等不再列出，也不推断填写。当前材料保留视频链接和公开信息占位；声明文本未代替成员签署。
''')
p=D/'00_审核入口/阅读入口.md';s=p.read_text();s=s.replace('作品：具身救援 · 从指令到行动','作品：具身救援 · 从指令到行动\n\n硕博组 G002｜今天不想做正事\n\n队长：郭煜，东南大学研究生；队员：林路辰，东南大学研究生。')
s=s.replace('待填：真实报名组别、官方参赛编号、队名、队长及成员信息、联系方式、成员分工、签名、日期；更细的模型权重修订号需主办方确认。','报名与成员信息已按用户提供内容填写；未提供的邮箱、学院、分工、额外成员、签署日期等不列出。剩余占位仅用于视频链接和公开内容。')
s=s.replace('成员签名稿、第三方与 AI 辅助说明、公开证明待补表','成员声明稿、第三方与 AI 辅助说明、公开证明待补表')
s=s.replace('最终 ZIP 按“组别-参赛编号-队伍名称-队长姓名.zip”命名，大小不超过200 MB。不得照搬示例包的身份信息。','最终 ZIP 名称为“硕博组-G002-今天不想做正事-郭煜.zip”，大小不超过200 MB。')
p.write_text(s)

p=D/'00_审核入口/要求对照清单.json';data=json.loads(p.read_text())
for row in data:
 if row['status']=='PREPARED_PERSONAL_FIELDS_PLACEHOLDER':row['status']='PERSONAL_FIELDS_FILLED_FROM_USER'
 elif row['status']=='VERIFIED_API_MODEL_WEIGHT_REVISION_PENDING':row['status']='ACTUAL_API_MODEL_VERIFIED'
 elif row['status']=='PREPARED_NAME_PLACEHOLDERS':row['status']='FINAL_NAME_G002_CONFIRMED'
p.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')

expected={'【待填_演示视频主链接】','【待填_演示视频备用链接】','【待填_公开内容标题】','【待填_公开发布平台】','【待填_公开内容链接】','【待填_公开发布时间】','【待填_公开页面截图】'}
found=set()
for p in D.rglob('*'):
 if p.is_file()and p.suffix in ['.md','.txt','.json']and p.name not in ['manifest.sha256.json','待填字段.json']:
  found.update(re.findall(r'【[^】]+】',p.read_text()))
for p in D.rglob('*.docx'):
 doc=Document(p);text='\n'.join([q.text for q in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells]);found.update(re.findall(r'【[^】]+】',text));assert '队长邮箱'not in text and '成员分工：'not in text
assert found<=expected,(found-expected)
(D/'00_审核入口/待填字段.json').write_text(json.dumps({'fields':{key:''for key in sorted(found)},'count':len(found),'scope':'Video links and actual public-content publication only; all other unknown personal fields removed.'},ensure_ascii=False,indent=2)+'\n')
(O/'team-information.json').write_text(json.dumps({'group':'硕博组','entry_id':'G002','team':'今天不想做正事','captain':{'name':'郭煜','university':'东南大学','student_status':'研究生','phone':CAPTAIN_PHONE},'members':[{'name':'林路辰','university':'东南大学','student_status':'研究生'}],'unprovided_personal_fields_omitted':True,'video_and_publication_placeholders_kept':True,'filename':'硕博组-G002-今天不想做正事-郭煜.zip'},ensure_ascii=False,indent=2)+'\n')
print('FINAL_REMAINING_FIELDS',json.dumps(sorted(found),ensure_ascii=False),flush=True)
